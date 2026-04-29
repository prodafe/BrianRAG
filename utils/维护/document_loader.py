import os
import re
import base64
import fitz
import ollama
import logging
import pandas as pd
import io
from langchain_classic.schema import Document
from langchain_community.document_loaders import (
    TextLoader,
    PyPDFLoader,
    UnstructuredWordDocumentLoader,
    UnstructuredHTMLLoader,
    CSVLoader,
    UnstructuredMarkdownLoader,
)
from config import Config
from docx import Document as DocxDocument  # python-docx

logger = logging.getLogger(__name__)

# ==================== 图片描述生成函数 ====================
def generate_image_caption(image_path: str) -> str:
    try:
        with open(image_path, "rb") as img_file:
            img_base64 = base64.b64encode(img_file.read()).decode('utf-8')
        response = ollama.chat(
            model=Config.VISION_MODEL,
            messages=[{
                'role': 'user',
                'content': '你是一个图片描述助手。请用中文详细描述这张图片的内容，只输出中文描述，不要包含英文。'
            }],
            images=[img_base64]
        )
        return response['message']['content']
    except Exception as e:
        logger.error(f"生成图片描述失败: {e}")
        return f"[图片: {os.path.basename(image_path)}]"

# ==================== 辅助函数 ====================
def load_markdown_with_images(file_path: str):
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()
    img_pattern = r'!\[.*?\]\((.*?)\)'
    images = re.findall(img_pattern, content)
    text_without_images = re.sub(img_pattern, '', content)
    doc = Document(page_content=text_without_images, metadata={"source": file_path, "images": images})
    return [doc]

def extract_images_from_markdown(content: str, base_dir: str) -> tuple[str, list]:
    img_pattern = r'!\[.*?\]\((.*?)\)'
    images = re.findall(img_pattern, content)
    abs_images = []
    for img in images:
        if not os.path.isabs(img):
            img = os.path.join(base_dir, img)
        abs_images.append(img)
    return content, abs_images

def extract_images_from_pdf(pdf_path: str) -> list:
    images = []
    doc = fitz.open(pdf_path)
    img_dir = Config.IMAGES_DIR
    os.makedirs(img_dir, exist_ok=True)
    for page_num in range(len(doc)):
        page = doc[page_num]
        img_list = page.get_images(full=True)
        for img_index, img in enumerate(img_list):
            xref = img[0]
            pix = fitz.Pixmap(doc, xref)
            if pix.n - pix.alpha < 4:
                pix = fitz.Pixmap(fitz.csRGB, pix)
            img_filename = f"{os.path.basename(pdf_path)}_p{page_num+1}_img{img_index}.png"
            img_path = os.path.join(img_dir, img_filename)
            pix.save(img_path)
            pix = None
            images.append((page_num, img_path))
    doc.close()
    return images

def extract_images_from_docx(docx_path: str) -> list:
    images = []
    doc = DocxDocument(docx_path)
    img_dir = Config.IMAGES_DIR
    os.makedirs(img_dir, exist_ok=True)
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            img_blob = rel.target_part.blob
            content_type = rel.target_part.content_type
            if "jpeg" in content_type:
                ext = "jpg"
            elif "png" in content_type:
                ext = "png"
            elif "gif" in content_type:
                ext = "gif"
            elif "bmp" in content_type:
                ext = "bmp"
            else:
                ext = "img"
            img_filename = f"{os.path.basename(docx_path)}_img_{len(images)}.{ext}"
            img_path = os.path.join(img_dir, img_filename)
            with open(img_path, "wb") as f:
                f.write(img_blob)
            images.append(img_path)
    return images

def extract_images_from_html(html_path: str) -> list:
    from bs4 import BeautifulSoup
    images = []
    with open(html_path, 'r', encoding='utf-8') as f:
        soup = BeautifulSoup(f, 'html.parser')
    for img in soup.find_all('img'):
        src = img.get('src')
        if src:
            if not os.path.isabs(src):
                src = os.path.join(os.path.dirname(html_path), src)
            if os.path.exists(src):
                images.append(src)
    return images

# ==================== 主加载函数 ====================
def load_single_document(file_path: str):
    ext = os.path.splitext(file_path)[1].lower()
    images = []
    if ext == ".txt":
        loader = TextLoader(file_path, encoding="utf-8")
        docs = loader.load()
    elif ext == ".pdf":
        img_infos = extract_images_from_pdf(file_path)
        loader = PyPDFLoader(file_path)
        docs = loader.load()
        for page_num, img_path in img_infos:
            caption = generate_image_caption(img_path)
            img_doc = Document(
                page_content=f"[图片描述] {caption}",
                metadata={
                    "source": file_path,
                    "page": page_num + 1,
                    "image_path": img_path,
                    "image_caption": caption,
                    "images": [img_path]
                }
            )
            docs.append(img_doc)
    elif ext == ".docx":
        images = extract_images_from_docx(file_path)
        loader = UnstructuredWordDocumentLoader(file_path)
        docs = loader.load()
    elif ext in [".html", ".htm"]:
        images = extract_images_from_html(file_path)
        loader = UnstructuredHTMLLoader(file_path)
        docs = loader.load()
    elif ext == ".md":
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()

        # 提取表格
        table_pattern = r'(\|.*\|\n\|[-: |]+\|\n(?:\|.*\|\n)+)'
        tables = re.findall(table_pattern, content, re.MULTILINE)
        text_without_tables = re.sub(table_pattern, '', content, flags=re.MULTILINE)

        docs = []
        # 纯文本
        if text_without_tables.strip():
            docs.append(Document(page_content=text_without_tables, metadata={"source": file_path, "type": "text"}))

        # 表格处理：保留原始表格 + 每行转换为键值对句子
        for table_idx, table_md in enumerate(tables):
            # 保留原始表格块（type: table_raw）
            docs.append(Document(page_content=table_md,
                                 metadata={"source": file_path, "type": "table_raw", "table_id": table_idx}))

            # 转换为键值对行
            try:
                import pandas as pd
                import io
                df = pd.read_html(io.StringIO(table_md))[0]
                headers = df.columns.tolist()
                for row_idx, row in df.iterrows():
                    row_parts = []
                    for col in headers:
                        val = row[col]
                        if pd.notna(val) and str(val).strip():
                            row_parts.append(f"{col}: {val}")
                    if row_parts:
                        row_desc = "，".join(row_parts)
                        docs.append(Document(
                            page_content=row_desc,
                            metadata={
                                "source": file_path,
                                "type": "table_row",
                                "row_index": row_idx,
                                "table_id": table_idx  # 关联到原始表格的 table_id
                            }
                        ))
            except Exception as e:
                logger.warning(f"表格行转换失败: {e}")

        # 图片描述（原有逻辑）
        _, images = extract_images_from_markdown(content, os.path.dirname(file_path))
        for img_path in images:
            if os.path.exists(img_path):
                caption = generate_image_caption(img_path)
                docs.append(Document(
                    page_content=f"[图片描述] {caption}",
                    metadata={"source": file_path, "image_path": img_path, "type": "image"}
                ))

    elif ext == ".csv":
        loader = CSVLoader(file_path)
        docs = loader.load()
    elif ext in [".jpg", ".jpeg", ".png", ".gif", ".bmp"]:
        caption = generate_image_caption(file_path)
        doc = Document(page_content=caption, metadata={"source": file_path, "images": [file_path]})
        docs = [doc]
    else:
        raise ValueError(f"不支持的文件类型: {ext}")

    # 将提取的图片路径存入每个 Document 的 metadata
    for doc in docs:
        if "images" not in doc.metadata:
            doc.metadata["images"] = []
        doc.metadata["images"].extend(images)
    return docs